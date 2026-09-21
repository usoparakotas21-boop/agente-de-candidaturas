import unittest
from io import BytesIO

from docx import Document
from pypdf import PdfReader

from app.document_pdf import docx_to_pdf


class DocumentPdfConversionTests(unittest.TestCase):
    def test_converts_generated_docx_to_readable_pdf_with_accents(self):
        source = Document()
        source.add_heading("Currículo de Analista de RH", level=0)
        source.add_paragraph("Experiência com recrutamento — seleção • comunicação…")
        content = BytesIO()
        source.save(content)

        pdf = docx_to_pdf(content.getvalue(), title="Currículo direcionado")
        self.assertTrue(pdf.startswith(b"%PDF-"))
        reader = PdfReader(BytesIO(pdf))
        self.assertEqual(len(reader.pages), 1)
        extracted = reader.pages[0].extract_text()
        self.assertIn("Currículo de Analista de RH", extracted)
        self.assertIn("recrutamento", extracted)
        self.assertIn("seleção", extracted)
        self.assertNotIn("?", extracted)


if __name__ == "__main__":
    unittest.main()
